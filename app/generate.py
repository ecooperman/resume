"""
Self-service resume.yaml generation from uploaded source material: a
LinkedIn profile export (see the Chrome extension's "Export LinkedIn
profile" button, job-application-chrome-extension/content/exportProfileContent.js)
plus any other documents someone has lying around (old resumes as PDF/DOCX,
plain text, LinkedIn's own "Get a copy of your data" CSV export, etc.) and
freeform notes.

This never saves anything itself - app/admin.py's /api/generate-resume
route returns the generated YAML text to the admin UI for a human to review
and edit before it becomes an actual Person row (see admin_static/app.js).
That review step is the whole reason this lives here as "generate text" and
not "generate and save a Person" - a first draft synthesized by Claude from
messy source documents should always get a human's eyes on it first.
"""
import io
import re
from pathlib import Path
from typing import List, Optional, Tuple

import anthropic
import yaml
from pypdf import PdfReader
from docx import Document

from . import config

ROOT = Path(__file__).resolve().parent.parent
EXAMPLE_RESUME_PATH = ROOT / "example-resume.yaml"

MAX_DOCUMENT_CHARS = 60_000  # per document - keeps one huge/garbled upload from blowing the whole prompt's budget


class ResumeGenerationError(Exception):
    """Missing config, an unreadable upload, or a Claude response that never
    became valid resume.yaml even after one repair attempt. Always has a
    message safe to show directly in the admin UI."""


def extract_text(filename: str, content: bytes) -> str:
    """Best-effort plain text out of one uploaded file, based on its
    extension. Anything not recognized (.txt, .md, .csv, .json, .yaml/.yml -
    including LinkedIn's own data-export CSVs and the extension's profile
    export JSON) is just decoded as text; that covers every source document
    this feature expects other than PDF/DOCX."""
    ext = Path(filename).suffix.lower()
    try:
        if ext == ".pdf":
            reader = PdfReader(io.BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif ext == ".docx":
            doc = Document(io.BytesIO(content))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        else:
            text = content.decode("utf-8", errors="ignore")
    except Exception as e:
        raise ResumeGenerationError(f"Couldn't read {filename!r}: {e}")
    text = text.strip()
    if len(text) > MAX_DOCUMENT_CHARS:
        text = text[:MAX_DOCUMENT_CHARS] + "\n...(truncated)"
    return text


SYSTEM_PROMPT_TEMPLATE = """You turn someone's source career material into a resume.yaml file for a
specific YAML schema. You will be given: a LinkedIn profile export (raw text
pulled from their profile page, not perfectly structured), some number of
other documents (old resumes, LinkedIn's own data-export CSVs, client lists,
freeform notes - also not perfectly structured), and possibly a short note
of extra context from the person themselves.

Rules:
- Never invent employers, titles, dates, schools, or accomplishments that
  aren't grounded in the source material. If a field genuinely isn't
  present anywhere, omit it (most fields are optional - see the schema)
  rather than guessing or writing a placeholder.
- Be comprehensive, not resume-length. Include every job, every bullet-
  worthy accomplishment, every skill, every certification you can find
  across all the sources - even ones a person might normally trim off a
  one-page resume. This file is meant to be a superset that later gets
  tailored down per job, not a finished printed resume, so err heavily
  toward including more rather than less.
- When the same job or fact appears in multiple sources with different
  wording, merge them into the single richest version rather than picking
  one source over another or duplicating the entry.
- Write real bullets, not just a copy of a LinkedIn one-line description -
  where the source material gives enough detail (metrics, scope, tools),
  turn it into a strong resume-style accomplishment bullet. Where it
  doesn't, keep it modest and factual rather than embellishing.
- Output ONLY the YAML content, no markdown code fences, no commentary
  before or after it. It must parse as YAML and have a top-level "basics"
  key.

Here is the exact schema, as a fully-populated fictional example - match
this shape exactly (same keys, same nesting):

{example_resume}
"""


def _build_system_prompt() -> str:
    example = EXAMPLE_RESUME_PATH.read_text()
    return SYSTEM_PROMPT_TEMPLATE.format(example_resume=example)


def _build_user_message(linkedin_text: Optional[str], documents: List[Tuple[str, str]], notes: Optional[str]) -> str:
    parts = []
    if linkedin_text:
        parts.append("=== LinkedIn profile export ===\n" + linkedin_text)
    for filename, text in documents:
        parts.append(f"=== Document: {filename} ===\n" + text)
    if notes and notes.strip():
        parts.append("=== Extra notes from the person ===\n" + notes.strip())
    if not parts:
        raise ResumeGenerationError("No source material provided - upload a LinkedIn export and/or a document first.")
    return "\n\n".join(parts)


def _extract_yaml(text: str) -> str:
    """Strip a ```yaml ... ``` fence if the model wrapped the output in one
    despite being asked not to - same cheap-insurance pattern as
    time-management's resume_gen.py._extract_json."""
    text = text.strip()
    match = re.match(r"^```(?:ya?ml)?\s*(.*?)\s*```$", text, re.DOTALL)
    return match.group(1) if match else text


def generate_resume_yaml(
    linkedin_text: Optional[str],
    documents: List[Tuple[str, str]],
    notes: Optional[str],
) -> str:
    """Calls Claude to synthesize a resume.yaml from the given source
    material. Returns the YAML text (already validated to parse and have a
    top-level "basics" key) - never saves anything. One repair attempt if
    the first response isn't valid before giving up."""
    if not config.ANTHROPIC_API_KEY:
        raise ResumeGenerationError(
            "ANTHROPIC_API_KEY is not set - resume generation is unavailable. See README.md for setup."
        )
    client = anthropic.Anthropic(api_key=config.ANTHROPIC_API_KEY)
    system = _build_system_prompt()
    messages = [{"role": "user", "content": _build_user_message(linkedin_text, documents, notes)}]

    last_error = None
    for attempt in range(2):
        response = client.messages.create(
            model="claude-opus-5",
            max_tokens=8000,
            system=system,
            messages=messages,
        )
        raw = "".join(block.text for block in response.content if block.type == "text")
        candidate = _extract_yaml(raw)
        try:
            parsed = yaml.safe_load(candidate)
            if not isinstance(parsed, dict) or "basics" not in parsed:
                raise ValueError('missing a top-level "basics" key')
            return candidate
        except Exception as e:
            last_error = e
            messages.append({"role": "assistant", "content": raw})
            messages.append(
                {
                    "role": "user",
                    "content": (
                        f"That wasn't valid: {e}. Respond again with ONLY the corrected YAML, "
                        "no markdown fences, no commentary."
                    ),
                }
            )
    raise ResumeGenerationError(f"Claude's response never became valid resume YAML: {last_error}")
